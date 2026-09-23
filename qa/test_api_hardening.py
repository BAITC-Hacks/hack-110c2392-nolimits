"""Regression cases use api_factory's disposable SQLite, never a running server."""
from concurrent.futures import ThreadPoolExecutor, TimeoutError
from io import BytesIO
from threading import Event

import pandas as pd
from openpyxl import load_workbook


def operation(kind, quantity, key):
    return {'kind': kind, 'date': '2026-09-23', 'warehouse': 'ASTANA', 'partner': 'CUSTOMER',
            'client_request_id': key, 'lines': [{'sku': 'QA-001', 'product_name': 'Test cable',
                                                'category': 'Cables', 'quantity': quantity, 'unit_price': 100}]}


def balance(client):
    return next(row['current_stock'] for row in client.get('/api/inventory/stock').json()['rows']
                if row['sku'] == 'QA-001' and row['warehouse'] == 'ASTANA')


def test_empty_sales_restores_real_stock_instead_of_loading_demo(api_factory, monkeypatch):
    from app import main
    with api_factory() as client:
        assert client.post('/api/inventory/movements', json=operation('ADJUSTMENT', 7, 'initial')).status_code == 200
        response = client.post('/api/data/upload/sales', files={
            'file': ('empty.csv', b'date,sku,product_name,category,quantity,price,customer_id,warehouse\n', 'text/csv'),
        })
        assert response.status_code == 200
        assert balance(client) == 7
        stable_id = client.get('/api/editor/stock').json()['rows'][0]['row_id']
    monkeypatch.setattr(main, 'state', main.AppState())
    with api_factory() as client:
        assert balance(client) == 7
        assert client.get('/api/data/status').json()['datasets']['sales'] == 0
        assert client.get('/api/editor/stock').json()['rows'][0]['row_id'] == stable_id
        assert client.get('/api/inventory/movements').json()['total'] == 1


def test_editor_update_keeps_same_record_after_earlier_row_is_deleted(client):
    rows = client.get('/api/editor/sales').json()['rows']
    target_id = rows[1]['row_id']
    original_third = rows[2]
    payload = {key: value for key, value in rows[1].items() if key != 'row_id'}
    payload['quantity'] = 99
    assert client.delete(f"/api/editor/sales/rows/{rows[0]['row_id']}").status_code == 200
    assert client.patch(f'/api/editor/sales/rows/{target_id}', json={'row': payload}).status_code == 200
    saved = client.get('/api/editor/sales').json()['rows']
    assert saved[0]['row_id'] == target_id and saved[0]['quantity'] == 99
    assert saved[1] == original_third
    assert client.patch(f"/api/editor/sales/rows/{rows[0]['row_id']}", json={'row': payload}).status_code == 404


def test_replacement_upload_invalidates_old_editor_identity(client):
    old_id = client.get('/api/editor/stock').json()['rows'][0]['row_id']
    replacement = b'sku,warehouse,current_stock\nQA-001,ASTANA,20\n'
    assert client.post('/api/data/upload/stock', files={'file': ('stock.csv', replacement, 'text/csv')}).status_code == 200
    new_id = client.get('/api/editor/stock').json()['rows'][0]['row_id']
    assert isinstance(new_id, int) and 0 <= new_id < 2 ** 53
    assert new_id != old_id
    assert client.patch(f'/api/editor/stock/rows/{old_id}', json={'row': {'current_stock': 999}}).status_code == 404
    assert balance(client) == 20


def test_catalog_delete_archives_and_survives_restart(api_factory, monkeypatch):
    from app import main
    with api_factory() as client:
        product = {'sku': 'ARCHIVE-ME', 'product_name': 'Temporary', 'category': 'Other', 'unit_price': 10, 'active': True}
        assert client.post('/api/editor/products/rows', json={'row': product}).status_code == 200
        row = next(r for r in client.get('/api/editor/products').json()['rows'] if r['sku'] == product['sku'])
        assert client.delete(f"/api/editor/products/rows/{row['row_id']}").status_code == 200
        assert product['sku'] not in {r['sku'] for r in client.get('/api/inventory/catalog').json()['products']}
        assert product['sku'] not in {r['sku'] for r in client.get('/api/editor/products').json()['rows']}
        assert client.patch(f"/api/editor/products/rows/{row['row_id']}", json={'row': {'product_name': 'Stale form'}}).status_code == 404
    monkeypatch.setattr(main, 'state', main.AppState())
    with api_factory() as client:
        assert product['sku'] not in {r['sku'] for r in client.get('/api/inventory/catalog').json()['products']}
        archived = next(r for r in client.get('/api/inventory/catalog?include_inactive=true').json()['products'] if r['sku'] == product['sku'])
        assert archived['active'] is False


def test_editor_exports_escape_formulas_and_hide_internal_identity(client):
    row = client.get('/api/editor/products').json()['rows'][0]
    assert client.patch(f"/api/editor/products/rows/{row['row_id']}", json={'row': {'product_name': '=1+1'}}).status_code == 200
    response = client.get('/api/editor/products/export?format=xlsx')
    book = load_workbook(BytesIO(response.content), data_only=False)
    assert all(cell.data_type != 'f' for cells in book.active for cell in cells)
    assert '__row_id' not in [cell.value for cell in book.active[1]]
    csv = client.get('/api/editor/products/export?format=csv').content.decode('utf-8-sig')
    assert "'=1+1" in csv and '__row_id' not in csv
    assert '__row_id' not in client.get('/api/editor/products').json()['columns']
    exported = pd.ExcelFile(BytesIO(client.get('/api/inventory/export').content))
    assert all('__row_id' not in exported.parse(sheet).columns for sheet in exported.sheet_names)


def test_idempotency_replays_same_document_and_rejects_different_payload(client):
    first = operation('ADJUSTMENT', 1, 'same-request')
    assert client.post('/api/inventory/movements', json=first).status_code == 200
    replay = client.post('/api/inventory/movements', json=first)
    assert replay.status_code == 200 and replay.json()['replayed'] is True
    changed = operation('ADJUSTMENT', 5, 'same-request')
    assert client.post('/api/inventory/movements', json=changed).status_code == 409
    assert balance(client) == 1
    assert client.get('/api/inventory/movements').json()['total'] == 1


def test_movement_and_editor_cannot_overwrite_each_others_snapshots(client, monkeypatch):
    from app import main
    row_id = client.get('/api/editor/stock').json()['rows'][0]['row_id']
    assert client.patch(f'/api/editor/stock/rows/{row_id}', json={'row': {'current_stock': 10}}).status_code == 200
    calculate = main.calculate_recommendations
    captured, resume = Event(), Event()

    def paused(data, *args, **kwargs):
        if float(data['stock'].iloc[0]['current_stock']) == 9:
            captured.set()
            assert resume.wait(10), 'Timed out waiting for the concurrent test'
        return calculate(data, *args, **kwargs)

    monkeypatch.setattr(main, 'calculate_recommendations', paused)
    with ThreadPoolExecutor(max_workers=2) as pool:
        sale = pool.submit(client.post, '/api/inventory/movements', json=operation('SALE', 1, 'serial-sale'))
        assert captured.wait(10)
        edit = pool.submit(client.patch, f'/api/editor/stock/rows/{row_id}', json={'row': {'current_stock': 20}})
        try:
            edit.result(timeout=0.1)
            raise AssertionError('The editor bypassed the shared snapshot lock')
        except TimeoutError:
            pass
        finally:
            resume.set()
        assert sale.result(timeout=10).status_code == 200
        assert edit.result(timeout=10).status_code == 200
    assert balance(client) == 20
    assert client.get('/api/inventory/movements').json()['total'] == 1


def test_failed_recommendation_save_rolls_back_entire_movement(client, monkeypatch):
    from app.repositories import database

    def fail(*_):
        raise RuntimeError('Simulated persistence failure')

    monkeypatch.setattr(database, '_save_recommendations', fail)
    response = client.post('/api/inventory/movements', json=operation('ADJUSTMENT', 7, 'failed'))
    assert response.status_code == 500
    assert balance(client) == 0
    assert client.get('/api/inventory/movements').json()['total'] == 0
    assert float(database.load_datasets_draft()['stock'].iloc[0]['current_stock']) == 0


def test_stale_balance_blocks_approval_until_explicit_stock_confirmation(client):
    from app import main
    main.state.datasets['stock']['balance_is_stale'] = True
    main.state.datasets['stock']['balance_unknown'] = True
    main.state.datasets['stock']['source_warning'] = 'Monthly opening snapshot'
    main.state.datasets['stock']['balance_date'] = pd.Timestamp('2026-03-01')
    main.state.datasets['stock']['as_of'] = pd.Timestamp('2026-03-25')
    assert client.post('/api/recommendations/calculate', json={}).status_code == 200
    recommendation = client.get('/api/recommendations').json()['recommendations'][0]
    assert recommendation['metadata']['balance_is_stale'] is True
    assert client.post(f"/api/orders/{recommendation['id']}/approve").status_code == 422
    row_id = client.get('/api/editor/stock').json()['rows'][0]['row_id']
    response = client.patch(f'/api/editor/stock/rows/{row_id}', json={'row': {'current_stock': 17}})
    assert response.status_code == 200, response.text
    saved = client.get('/api/editor/stock').json()['rows'][0]
    assert saved['balance_is_stale'] is False
    assert saved['balance_unknown'] is False
    assert saved['source_warning'] == ''
    assert saved['balance_date'].startswith('2026-03-25')
    assert client.post(f"/api/orders/{recommendation['id']}/approve").status_code == 200


def test_analytics_uses_same_global_cutoff_as_recommendation(client):
    from app import main
    later = main.state.datasets['sales'].iloc[-1:].copy()
    later['date'] = pd.Timestamp('2026-04-15')
    later['warehouse'] = 'ALMATY'
    main.state.datasets['sales'] = pd.concat([main.state.datasets['sales'], later], ignore_index=True)
    assert client.post('/api/recommendations/calculate', json={'safety_days': 3}).status_code == 200
    row = next(r for r in client.get('/api/recommendations').json()['recommendations'] if r['warehouse'] == 'ASTANA')
    analytics = client.get('/api/analytics/QA-001?warehouse=ASTANA').json()
    actual = [point for point in analytics['points'] if point.get('forecast') is None]
    forecast = [point for point in analytics['points'] if point.get('forecast') is not None]
    assert row['metadata']['as_of'] == '2026-04-15'
    assert actual[-1]['date'] == '2026-04-15'
    assert forecast[0]['date'] == '2026-04-16'
    assert len(forecast) == row['lead_time_days'] + 3
    assert abs(sum(p['forecast'] for p in forecast[:row['lead_time_days']]) - row['forecast_lead_time']) < 0.2


def test_partner_upload_requires_explicit_delivery_policy(client):
    before = client.get('/api/data/status').json()
    response = client.post('/api/data/upload-partner', files={'file': ('partner.zip', b'zip', 'application/zip')})
    assert response.status_code == 422
    assert client.get('/api/data/status').json() == before


def test_explicit_supplier_price_and_pack_replace_unknown_placeholders(client):
    from app import main
    main.state.datasets['suppliers']['unit_cost'] = 0.0
    main.state.datasets['suppliers']['cost_unknown'] = True
    main.state.datasets['suppliers']['package_size'] = 1
    main.state.datasets['suppliers']['package_unknown'] = True
    row_id = client.get('/api/editor/suppliers').json()['rows'][0]['row_id']
    url = f'/api/editor/suppliers/rows/{row_id}'
    assert client.patch(url, json={'row': {'lead_time_days': 12, 'unit_cost': 0, 'package_size': 1}}).status_code == 200
    unchanged = client.get('/api/editor/suppliers').json()['rows'][0]
    assert unchanged['cost_unknown'] is True and unchanged['package_unknown'] is True
    assert client.patch(url, json={'row': {'unit_cost': 250, 'package_size': 6}}).status_code == 200
    confirmed = client.get('/api/editor/suppliers').json()['rows'][0]
    assert confirmed['cost_unknown'] is False and confirmed['package_unknown'] is False
    assert client.get('/api/recommendations').json()['recommendations'][0]['metadata']['cost_unknown'] is False
